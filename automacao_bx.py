import pyautogui
import time
import pyperclip
import openpyxl

# =================================================================
# CONFIGURAÇÕES
# =================================================================
pyautogui.PAUSE = 0.8  
CONFIANCA = 0.8        

ARQUIVO_EXCEL = 'ECF 2026 (Ano calendário 2025).xlsx' 
ABAS_ALVO = ["SAIRAM", "Ativas_e_Inativas 130326"]
COLUNA_CNPJ = 'A'   
COLUNA_STATUS = 'J' 
LINHA_INICIAL = 2   

DATA_INICIO = "01012025" 
DATA_FIM = "31122025"    

# =================================================================
# FUNÇÕES
# =================================================================
def clicar(imagem, aguardar=10):
    inicio = time.time()
    while time.time() - inicio < aguardar:
        try:
            posicao = pyautogui.locateCenterOnScreen(imagem, confidence=CONFIANCA)
            if posicao:
                pyautogui.click(posicao)
                return True
        except:
            pass
        time.sleep(0.3)
    return False

def consultar_cnpj(cnpj):
    print(f"\n🔄 Iniciando troca para CNPJ: {cnpj}")
    
    # 1. Clica no ícone para abrir a janelinha
    if not clicar('trocar_perfil_icone.png'): return "ERRO IMAGEM"
    
    # 2. Preenche a janelinha de Perfil (Procurador + CNPJ)
    clicar('radio_procurador.png')
    clicar('campo_cnpj.png')
    
    pyautogui.hotkey('ctrl', 'a') # Seleciona o CNPJ antigo se houver
    pyautogui.press('backspace')
    pyperclip.copy(str(cnpj))
    pyautogui.hotkey('ctrl', 'v') 
    
    # 3. Confirma a troca no botão da janela pop-up
    clicar('botao_confirmar_perfil.png')
    
    print("Aguardando carregamento do perfil...")
    time.sleep(4) # Tempo para o sistema processar a troca de cliente

    # 4. Configura os filtros de pesquisa
    clicar('drop_sistema.png')
    clicar('opcao_sped.png')
    clicar('drop_arquivo.png')
    clicar('opcao_ecd.png')

    # 5. Datas
    clicar('prencher data inicial.png')
    pyautogui.write(DATA_INICIO)
    clicar('prencher data final.png')
    pyautogui.write(DATA_FIM)

    # 6. Busca
    clicar('botao_pesquisar.png')
    time.sleep(5) 

    # 7. Verifica Resultado
    try:
        if pyautogui.locateOnScreen('aviso_erro.png', confidence=CONFIANCA):
            clicar('botao_ok_erro.png')
            return "NÃO ENCONTRADO"
    except:
        pass
    
    return "SIM"

# =================================================================
# EXECUÇÃO
# =================================================================
wb = openpyxl.load_workbook(ARQUIVO_EXCEL)
print("Selecione o ReceitaNet BX em 5 segundos...")
time.sleep(5)

for aba_nome in ABAS_ALVO:
    if aba_nome in wb.sheetnames:
        planilha = wb[aba_nome]
        for linha in range(LINHA_INICIAL, planilha.max_row + 1):
            cnpj = planilha[f"{COLUNA_CNPJ}{linha}"].value
            status = planilha[f"{COLUNA_STATUS}{linha}"].value
            
            if not cnpj or status: continue
                
            res = consultar_cnpj(cnpj)
            planilha[f"{COLUNA_STATUS}{linha}"] = res
            wb.save(ARQUIVO_EXCEL)
            print(f"✅ Linha {linha} finalizada: {res}")

print("\n🚀 Processo concluído!")